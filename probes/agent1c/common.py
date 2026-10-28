"""common.py — shared mock builders for agent1c probes (data-handling audit).

Builds tiny synthetic Excel trees matching Accurate export formats and a
config.json so probes can run REAL app functions (build_append1, fn_*, etc).
"""
import json
import os
import sys
import tempfile
from datetime import datetime

import openpyxl

PROJECT = "/home/z/stock-audit/stock-monitoring-python"
sys.path.insert(0, PROJECT)

SHEET_2205 = "#22.05. Rincian Penjualan per B"
SHEET_3224 = "#32.24. Rincian Penerimaan Bara"
SHEET_4205 = "#42.05. Rincian Pemindahan Bara"
SHEET_4215 = "#42.15. Pemindahan Barang (KIRI"
SHEET_4217 = "#42.17. Kuantitas Barang per Gu"

FOLDERS = {
    "RekapBerjalan":      "1. DEVIASI WEEKLY",
    "RekapLalu":          "1.12 Rekap Mentah Bulan Lalu",
    "Penjualan2205":      "2. 22.05",
    "Kirim4215":          "3.1 42.15 Kirim",
    "RI3224":             "13. 32.24 RI",
    "RI3224Lalu":         "13.2 3224 RI Bulan Lalu",
    "SaldoAwal4217":      "6.1 42.17 Saldo Awal",
    "SaldoSistem4217":    "6. 42.17",
    "DBSKategori":        "21. DBS Kategori",
    "Pemindahan4205":     "4. 42.05",
    "Pemindahan4205Lalu": "4.1 4205 Bulan Lalu",
    "Data9901":           "30. 9901",
}
SHEETS = {
    "RekapMentah":     "REKAP MENTAH",
    "Penjualan2205":   SHEET_2205,
    "Pemindahan4215":  SHEET_4215,
    "Penerimaan3224":  SHEET_3224,
    "SaldoAwal4217":   SHEET_4217,
    "DBSKategori":     "Sheet1",
    "Pemindahan4205":  SHEET_4205,
}


def write_xlsx(path, sheet_name, header, rows, header_row_idx=0):
    """Write one-sheet xlsx. header_row_idx = 0-based row of the header."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    for c, h in enumerate(header, 1):
        ws.cell(row=header_row_idx + 1, column=c, value=h)
    for r, row in enumerate(rows, header_row_idx + 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    wb.save(path)
    wb.close()
    return path


def make_root():
    d = tempfile.mkdtemp(prefix="a1c_")
    root = os.path.join(d, "ROOT")
    os.makedirs(root, exist_ok=True)
    return root


def make_config(root, extra=None):
    cfg = {"RootPath": root + os.sep, "Folders": dict(FOLDERS),
           "Sheets": dict(SHEETS)}
    if extra:
        for k, v in extra.items():
            if k == "Folders":
                cfg["Folders"].update(v)
            elif k == "Sheets":
                cfg["Sheets"].update(v)
            else:
                cfg[k] = v
    p = os.path.join(os.path.dirname(root.rstrip(os.sep)), "config.json")
    with open(p, "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2, ensure_ascii=False)
    return p


# ---------------------------------------------------------------- builders
def b_rekap(root, resto, barang_rows, mode="berjalan", filename=None):
    """Per-resto REKAP MENTAH file. barang_rows = [(nama, satuan, jenis,
    source, {day: val})] or + optional 6th element kategori (writes a
    KATEGORI column after NAMA RESTO, v3.9.49 schema)."""
    days = list(range(1, 32)) if mode == "berjalan" else list(range(25, 32))
    has_kat = any(len(r) >= 6 and r[5] is not None for r in barang_rows)
    header = ["NAMA RESTO"] + (["KATEGORI"] if has_kat else []) + \
        ["SOURCE DATA", "JENIS", "STATUS", "NAMA BARANG",
         "SATUAN"] + [str(d) for d in days]
    rows = []
    for r in barang_rows:
        nama, satuan, jenis, src, vals = r[:5]
        kat = r[5] if len(r) >= 6 else None
        rows.append([resto] + ([kat] if has_kat else []) +
                    [src, jenis, "OK", nama, satuan]
                    + [vals.get(d, None) for d in days])
    folder = "1. DEVIASI WEEKLY" if mode == "berjalan" \
        else "1.12 Rekap Mentah Bulan Lalu"
    return write_xlsx(os.path.join(root, folder, filename or f"{resto}.xlsx"),
                      "REKAP MENTAH", header, rows)


def b_2205(root, rows):
    """rows = [(nomor, tanggal, pelanggan, barang, qty)]"""
    header = ["Nomor #", "Tanggal", "Nama Pelanggan", "Nama Barang",
              "Kuantitas"]
    return write_xlsx(os.path.join(root, "2. 22.05", "penjualan.xlsx"),
                      SHEET_2205, header, rows, header_row_idx=4)


def b_3224(root, rows, mode="berjalan"):
    """rows = [(tanggal, kts, barang, gudang, nomor_ri, nomor_po, tgl_jam)]"""
    header = ["Tanggal", "Kts (Unit#1)", "Nama Barang", "Nama Gudang",
              "Nomor # RI", "Nomor # PO", "Tgl/Jam Pembuatan"]
    folder = "13. 32.24 RI" if mode == "berjalan" \
        else "13.2 3224 RI Bulan Lalu"
    return write_xlsx(os.path.join(root, folder, "ri.xlsx"), SHEET_3224,
                      header, rows, header_row_idx=4)


def b_4205(root, rows, mode="berjalan"):
    """rows = [(nomor, gudang, tanggal, barang, qty, tgl_jam)]"""
    header = ["Nomor #Terima", "Gudang #Terima", "Tanggal #Terima",
              "Nama Barang", "#Qty. Terkecil", "#Tgl/Jam Pembuatan IT"]
    folder = "4. 42.05" if mode == "berjalan" else "4.1 4205 Bulan Lalu"
    return write_xlsx(os.path.join(root, folder, "pemindahan.xlsx"),
                      SHEET_4205, header, rows, header_row_idx=4)


def b_4215(root, rows):
    """rows = [(tanggal, gudang_kirim, barang, kts)]"""
    header = ["Tanggal", "Gudang Pengirim", "Nama Barang", "Kts Diterima"]
    return write_xlsx(os.path.join(root, "3.1 42.15 Kirim", "kirim.xlsx"),
                      SHEET_4215, header, rows, header_row_idx=4)


def b_4217(root, rows, folder="6.1 42.17 Saldo Awal", fname="saldo.xlsx"):
    """rows = [(nama_barang, kategori, satuan, {gudang_col: value})]"""
    gudangs = []
    for r in rows:
        for g in r[3]:
            if g not in gudangs:
                gudangs.append(g)
    header = ["Nama Barang", "Kategori Barang", "Nama Satuan"] + gudangs
    body = [[r[0], r[1], r[2]] + [r[3].get(g) for g in gudangs] for r in rows]
    return write_xlsx(os.path.join(root, folder, fname), SHEET_4217,
                      header, body, header_row_idx=4)


def b_dbs(root, rows):
    """rows = [(nama_barang, kategori, aktif)]"""
    header = ["NAMA BARANG", "KATEGORI", "JENIS", "SOURCE KIRIMAN",
              "AKTIF/NON"]
    body = [[r[0], r[1], "X", "", r[2]] for r in rows]
    return write_xlsx(os.path.join(root, "21. DBS Kategori", "kategori.xlsx"),
                      "Sheet1", header, body)


def b_9901(root, rows, fname="mutasi.xlsx", sheet="Sheet1"):
    """rows = list of dicts (raw sheet rows) written in fixed column order."""
    cols = ["Nomor #", "Tanggal", "Nama barang", "Nama cabang",
            "#Prime.Qty", "Nomor # Penerimaan Barang",
            "Jenis Barang Barang & Jasa", "Nomor # Pesanan Pembelian"]
    header = cols
    body = []
    for r in rows:
        body.append([r.get(c) for c in cols])
    return write_xlsx(os.path.join(root, "30. 9901", fname), sheet,
                      header, body)


def long_subset(final, jenis, pair=None):
    """Rows of append1 long output for JENIS, optionally filtered by pair."""
    d = final[final["JENIS"] == jenis]
    if pair is not None:
        b, r = pair
        d = d[(d["Nama Barang"].astype(str) == str(b))
              & (d["NAMA RESTO"].astype(str) == str(r))]
    return d
