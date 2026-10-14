"""
helpers.py — Helper functions yang mirror persis fnBukaSheet, fnCombineFolder,
dan fungsi text delimiter di Power Query M code.

OPTIMIZED untuk speed: vectorized text operations via .str accessor.

v3.9.37 — CACHE BACAAN FILE SUMBER (user: "ada tombol refresh untuk file
sumber, bisa gak? dan gak harus refresh semua tapi"): fn_buka_sheet kini
menyimpan bacaan per (file, sheet, header_row, raw) divalidasi SIGNATURE
(mtime_ns + ukuran). File yang tak berubah dibaca SEKALI per sesi —
run berikutnya (termasuk Refresh Terpilih) menyajikan frame dari cache
(copy defensif — pemanggil bebas memutasi). Invalidasi selektif via
read_cache_invalidate(paths) (dipakai dialog 🔄 Refresh File Sumber);
hasil identik baca ulang — hanya lebih cepat. Batas memori: estimasi
sel object (guard _READ_CACHE_MAX_CELLS, evict FIFO).
"""
import glob
import os
import re
import threading
from typing import List, Optional

import numpy as np
import pandas as pd

# =============================================================================
# TEXT DELIMITER FUNCTIONS (vectorized via .str accessor)
# =============================================================================

def text_before_delimiter(s, delimiter: str = ".") -> Optional[str]:
    """Mirror M: Text.BeforeDelimiter(text, delimiter). Single value."""
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return None
    s = str(s)
    idx = s.find(delimiter)
    if idx < 0:
        return s
    return s[:idx]


def text_after_delimiter(s, delimiter: str = ".") -> Optional[str]:
    """Mirror M: Text.AfterDelimiter(text, delimiter). Single value."""
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return None
    s = str(s)
    idx = s.find(delimiter)
    if idx < 0:
        return s
    return s[idx + len(delimiter):]


def text_between_delimiters(s, open_delim: str = "(", close_delim: str = ")") -> Optional[str]:
    """Mirror M: Text.BetweenDelimiters(text, openDelim, closeDelim). Single value."""
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return None
    s = str(s)
    start = s.find(open_delim)
    if start < 0:
        return None
    start += len(open_delim)
    end = s.find(close_delim, start)
    if end < 0:
        return None
    return s[start:end]


# === VECTORIZED VERSIONS (for pandas Series — 10-50x faster than .apply) ===

def vec_text_before_delimiter(series: pd.Series, delimiter: str = ".") -> pd.Series:
    """Vectorized Text.BeforeDelimiter untuk pandas Series.

    Perilaku: baris mengandung delimiter → teks sebelum delimiter;
    tanpa delimiter → teks asli; NaN → None. (v3.3: dead lines dibersihkan
    — dulu `result` dihitung dua kali, hasil pertama ditimpa.)
    """
    s = series.astype(str).where(series.notna(), other=np.nan)
    idx = s.str.find(delimiter)
    result = s.copy()
    has_delim = idx >= 0
    if bool(has_delim.any()):
        positions = idx[has_delim].values
        vals = s[has_delim].values
        result[has_delim] = [str(v)[:int(p)] for v, p in zip(vals, positions)]
    result = result.where(series.notna(), other=None)
    return result


def vec_text_after_delimiter(series: pd.Series, delimiter: str = ".") -> pd.Series:
    """Vectorized Text.AfterDelimiter untuk pandas Series."""
    s = series.astype(str)
    parts = s.str.split(delimiter, n=1, expand=True)
    if parts.shape[1] < 2:
        # No delimiter found in any row
        return series
    result = parts[1].fillna(series.astype(str))
    result = result.where(series.notna(), other=None)
    return result


def vec_text_between_delimiters(series: pd.Series, open_delim: str = "(", close_delim: str = ")") -> pd.Series:
    """Vectorized Text.BetweenDelimiters untuk pandas Series.
    Uses regex: extract text between first open_delim and next close_delim.
    """
    # Escape regex special chars in delimiters
    open_esc = re.escape(open_delim)
    close_esc = re.escape(close_delim)
    pattern = f"{open_esc}([^{close_esc}]*){close_esc}"
    result = series.astype(str).str.extract(pattern, expand=False)
    # Where no match, result is NaN → convert to None
    result = result.where(result.notna(), other=None)
    # Also preserve original NaN
    result = result.where(series.notna(), other=None)
    return result


# =============================================================================
# COMBINE FOLDER (mirror fnCombineFolder + fnBukaSheet)
# =============================================================================

def _list_excel_files(folder_path: str) -> List[str]:
    """List semua file Excel di folder, exclude hidden & temp files.
    v3.3.1: sertakan *.xlsb — sumber SO user kini berupa workbook gabungan
    .xlsb ("01.1 CEK STOCK V2.xlsb") yang sebelumnya TIDAK terbaca sama
    sekali (akibatnya Saldo SO kosong semua di tab Deviasi GR/PCS)."""
    if not os.path.isdir(folder_path):
        return []
    patterns = ["*.xlsx", "*.xls", "*.xlsm", "*.xlsb"]
    files = []
    for p in patterns:
        files.extend(glob.glob(os.path.join(folder_path, p)))
    # Filter: exclude hidden files (check basename, not full path) & temp lock files (~$...)
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    files = [f for f in files if not os.path.basename(f).startswith(".")]
    return sorted(files)


def _get_excel_engine() -> str:
    """Detect best available Excel engine.
    Priority: calamine (Rust, 10-20x faster) > openpyxl (default).
    """
    try:
        import python_calamine  # noqa: F401
        return "calamine"
    except ImportError:
        return None  # default to openpyxl


# Cache engine detection (only check once)
_EXCEL_ENGINE = _get_excel_engine()


# =============================================================================
# v3.9.37 — CACHE BACAAN FILE SUMBER (refresh selektif tanpa baca ulang
# file yang tak berubah; user: "gak harus refresh semua")
# =============================================================================
_READ_CACHE: dict = {}          # key -> {"sig", "df", "cells"}
_READ_CACHE_LOCK = threading.Lock()
_READ_CACHE_CELLS = 0           # total sel ter-cache (guard memori)
_READ_CACHE_MAX_CELLS = 20_000_000   # ~estimasi 20 juta sel object
_READ_CACHE_HIT_FILES: set = set()   # normpath file yang disajikan cache
_READ_CACHE_MISS_FILES: set = set()  # normpath file yang dibaca disk


def _read_cache_key(file_path: str, sheet_name, header_row: int,
                    raw: bool):
    """Kunci cache: path ternormalisasi (Windows case-insensitive) +
    sheet + header_row + raw — signature file divalidasi TERPISAH."""
    return (os.path.normcase(os.path.normpath(file_path)),
            str(sheet_name), int(header_row), bool(raw))


def _read_cache_sig(file_path: str):
    """Signature file: (mtime_ns, ukuran). None = file tak terbaca stat
    (hilang/terkunci) → cache MISS (baca ulang; kegagalan mengikuti jalur
    error fn_buka_sheet yang sudah ada)."""
    try:
        st = os.stat(file_path)
        return (st.st_mtime_ns, st.st_size)
    except OSError:
        return None


def _read_cache_get(key, sig):
    """Ambil frame dari cache bila signature cocok → COPY defensif
    (pemanggil boleh memutasi hasilnya — insert kolom Name / ffill di
    fn_combine_folder — tanpa menodai entry cache)."""
    with _READ_CACHE_LOCK:
        entry = _READ_CACHE.get(key)
        if entry is None or sig is None or entry["sig"] != sig:
            return None
        _READ_CACHE_HIT_FILES.add(key[0])
        return entry["df"].copy()


def _read_cache_put(key, sig, df) -> None:
    """Simpan copy frame ke cache (evict FIFO bila guard sel terlampaui).
    Frame yang DISIMPAN adalah copy — frame asli dikembalikan ke
    pemanggil (boleh dimutasi bebas)."""
    if sig is None or df is None:
        return
    try:
        cells = int(df.shape[0]) * int(df.shape[1])
    except (TypeError, ValueError):
        return
    if cells <= 0:
        return
    global _READ_CACHE_CELLS
    with _READ_CACHE_LOCK:
        # evict FIFO sampai muat (dict = urutan insert)
        while (_READ_CACHE_CELLS + cells > _READ_CACHE_MAX_CELLS
               and _READ_CACHE):
            _old_k, _old = next(iter(_READ_CACHE.items()))
            _READ_CACHE.pop(_old_k, None)
            _READ_CACHE_CELLS -= _old.get("cells", 0)
        if _READ_CACHE_CELLS + cells > _READ_CACHE_MAX_CELLS:
            return  # file raksasa melebihi guard sendirian → tak di-cache
        _READ_CACHE[key] = {"sig": sig, "df": df.copy(), "cells": cells}
        _READ_CACHE_CELLS += cells
        _READ_CACHE_MISS_FILES.add(key[0])


def read_cache_invalidate(paths=None) -> int:
    """v3.9.37 — buang entry cache file sumber. Return jumlah entry
    yang dibuang.

    paths=None → SEMUA entry dibuang ("Refresh Semua" = baca ulang
    semua file dari disk). Bila list/set path → hanya entry yang
    file-nya termasuk (pencocokan normcase+normpath — jalur Windows
    beda kapital / separator tetap kena)."""
    global _READ_CACHE_CELLS
    with _READ_CACHE_LOCK:
        if paths is None:
            n = len(_READ_CACHE)
            _READ_CACHE.clear()
            _READ_CACHE_CELLS = 0
            return n
        targets = {os.path.normcase(os.path.normpath(str(p)))
                   for p in paths}
        drop = [k for k in _READ_CACHE if k[0] in targets]
        for k in drop:
            _READ_CACHE_CELLS -= _READ_CACHE[k].get("cells", 0)
            _READ_CACHE.pop(k, None)
        return len(drop)


def read_cache_stats() -> tuple:
    """v3.9.37 — (n_entry, n_file_unik, n_file_hit, n_file_miss) utk
    log observabilitas (file ter-cache / disajikan cache / dibaca disk)."""
    with _READ_CACHE_LOCK:
        files = {k[0] for k in _READ_CACHE}
        return (len(_READ_CACHE), len(files),
                len(_READ_CACHE_HIT_FILES), len(_READ_CACHE_MISS_FILES))


def read_cache_reset_counters() -> None:
    """v3.9.37 — reset penghitung hit/miss per-file (dipanggil app di
    AWAL tiap run pipeline supaya log akhir run mencerminkan run itu)."""
    with _READ_CACHE_LOCK:
        _READ_CACHE_HIT_FILES.clear()
        _READ_CACHE_MISS_FILES.clear()


def _xlsb_engine() -> Optional[str]:
    """Engine untuk file .xlsb: calamine (jika ada) > pyxlsb > None.
    None = engine tidak tersedia (.xlsb harus di-skip dengan pesan jelas)."""
    if _EXCEL_ENGINE == "calamine":
        return "calamine"
    try:
        import pyxlsb  # noqa: F401
        return "pyxlsb"
    except ImportError:
        return None


def _list_sheet_names(file_path: str) -> List[str]:
    """Daftar nama sheet dari file Excel (semua format, termasuk .xlsb).
    Urutan coba: calamine > openpyxl (.xlsx) / pyxlsb (.xlsb)."""
    low = file_path.lower()
    # 1) calamine — cepat untuk semua format
    if _EXCEL_ENGINE == "calamine":
        try:
            from python_calamine import load_workbook
            return list(load_workbook(file_path).sheet_names)
        except Exception:
            pass
    # 2) .xlsb → pyxlsb
    if low.endswith(".xlsb"):
        try:
            from pyxlsb import open_workbook
            with open_workbook(file_path) as wb:
                return list(wb.sheets)
        except Exception:
            return []
    # 3) .xlsx/.xlsm → openpyxl
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True)
        names = list(wb.sheetnames)
        wb.close()
        return names
    except Exception:
        return []


def fn_buka_sheet(file_path: str, sheet_name: str, header_row: int = 0,
                   raw: bool = False) -> pd.DataFrame:
    """
    Mirror M: fnBukaSheet(binary, sheetName, doPromote).
    Baca sheet tertentu dari file Excel.

    OPTIMIZED: gunakan calamine engine (Rust) jika tersedia — 10-20x faster.
    Fallback ke openpyxl jika calamine tidak terinstall.
    v3.3.1: file .xlsb didukung (engine pyxlsb bila calamine absen).

    raw=True: baca tanpa header (header=None), kolom = Column1, Column2, ...
              Mirror M doPromote=null/false.
    raw=False: baca dengan header di baris header_row (0-indexed).
              Mirror M doPromote=true + Skip(N) + PromoteHeaders.
              header_row=4 = skip rows 0-3, row 4 jadi header.
    """
    # v3.3.1: pilih engine eksplisit untuk .xlsb (openpyxl tidak bisa .xlsb;
    # pandas default-engine pun butuh pyxlsb terinstall — beri pesan jelas)
    # v3.9.37: cache bacaan (signature mtime+ukuran) — file tak berubah
    # tak dibaca ulang dari disk (Refresh Terpilih: hanya file pilihan
    # user yang dibaca ulang; hasil identik).
    _rc_key = _read_cache_key(file_path, sheet_name, header_row, raw)
    _rc_sig = _read_cache_sig(file_path)
    _rc_hit = _read_cache_get(_rc_key, _rc_sig)
    if _rc_hit is not None:
        return _rc_hit
    if file_path.lower().endswith(".xlsb") and _EXCEL_ENGINE != "calamine":
        xlsb_eng = _xlsb_engine()
        if xlsb_eng is None:
            raise ImportError(
                "File .xlsb membutuhkan package 'pyxlsb' atau "
                "'python-calamine' — jalankan: pip install pyxlsb")
        engine_kwargs = {"engine": xlsb_eng}
    else:
        engine_kwargs = {"engine": _EXCEL_ENGINE} if _EXCEL_ENGINE else {}
    try:
        if raw:
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None,
                                dtype=object, **engine_kwargs)
        else:
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=header_row,
                                dtype=object, **engine_kwargs)
        _read_cache_put(_rc_key, _rc_sig, df)
        return df
    except Exception as e:
        # Fallback to openpyxl if calamine fails (e.g., unsupported format)
        if _EXCEL_ENGINE:
            if raw:
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=None, dtype=object)
            else:
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=header_row, dtype=object)
            _read_cache_put(_rc_key, _rc_sig, df)
            return df
        raise e


def fn_combine_folder(folder_path: str, sheet_name: str, header_row: int = 0,
                       raw: bool = False, add_name_col: bool = False,
                       ffill_cols: "list[str] | None" = None) -> pd.DataFrame:
    """
    Mirror M: fnCombineFolder(folderPath, sheetName, doPromote).

    raw=True: baca tanpa header per file, gabungkan raw.
              Mirror M doPromote=null/false.
    raw=False: baca dengan header_row per file, gabungkan by column name.
              Mirror M doPromote=true.

    add_name_col: jika True, tambah kolom "Name" berisi nama file (basename).
                  Diperlukan untuk fnRekapMentah (NAMA RESTO dari file name).

    ffill_cols: v3.9.33 — AUDIT FIX (48-a): list nama kolom yang di-ffill
                PER FILE sebelum concat. Nilai ber-pola grup header (mis.
                kolom 'Kategori Barang' 42.17 — hanya terisi di baris
                pertama grup) harus di-ffill DALAM file yang sama; ffill
                pada frame hasil CONCAT membuat file berikutnya mewarisi
                nilai EKOR file sebelumnya (kategori lintas file bleed —
                uji 48-a: file-2 grup pertama mewarisi kategori terakhir
                file-1). Bisa berisi beberapa kandidat nama kolom — semua
                yang ada di file di-ffill. Default None = perilaku lama.
    """
    files = _list_excel_files(folder_path)
    if not files:
        return pd.DataFrame()

    dfs = []
    for f in files:
        try:
            if raw:
                df = fn_buka_sheet(f, sheet_name, raw=True)
            else:
                df = fn_buka_sheet(f, sheet_name, header_row=header_row)
            if df is not None and len(df) > 0:
                if add_name_col:
                    # Insert Name column at position 0
                    df.insert(0, "Name", os.path.basename(f))
                if ffill_cols:
                    # v3.9.33 — AUDIT FIX (48-a): ffill PER FILE sebelum
                    # concat (mirror pola per-file fn_4109) — grup kategori
                    # file berikutnya TIDAK mewarisi ekor file sebelumnya.
                    for c in ffill_cols:
                        if c in df.columns:
                            df[c] = df[c].ffill()
                dfs.append(df)
        except Exception as e:
            print(f"  [WARNING] Skip file {os.path.basename(f)}: {e}")

    if not dfs:
        return pd.DataFrame()

    if raw:
        # Raw mode: concat by position (columns are 0, 1, 2, ...)
        # Use ignore_index + sort=False for speed
        combined = pd.concat(dfs, ignore_index=True, sort=False)
    else:
        # Header mode: concat by column name
        combined = pd.concat(dfs, ignore_index=True, sort=False)
    return combined


def promote_headers(df: pd.DataFrame) -> pd.DataFrame:
    """
    Mirror M: Table.PromoteHeaders(table, [PromoteAllScalars = true]).
    Row 0 jadi header, drop row 0.
    Jika ada duplicate column names after promote, keep first.
    """
    if df.empty:
        return df

    # Extract row 0 as headers
    headers = df.iloc[0].tolist()

    # Drop row 0
    result = df.iloc[1:].reset_index(drop=True)

    # Set new headers
    # Handle None/NaN in headers → replace with "ColumnN"
    new_cols = []
    seen = {}
    for i, h in enumerate(headers):
        if h is None or (isinstance(h, float) and pd.isna(h)):
            h = f"Column{i+1}"
        h = str(h)
        # Handle duplicates: append _1, _2, etc.
        if h in seen:
            seen[h] += 1
            h = f"{h}_{seen[h]}"
        else:
            seen[h] = 0
        new_cols.append(h)

    result.columns = new_cols
    return result


def table_skip(df: pd.DataFrame, n: int) -> pd.DataFrame:
    """Mirror M: Table.Skip(table, n). Skip first n rows."""
    if df.empty or n <= 0:
        return df
    return df.iloc[n:].reset_index(drop=True)


# =============================================================================
# v3.9.33 — PRE-FLIGHT FOLDER DATA (kontrak — dipanggil app.py agent 49-c)
# =============================================================================

def preflight_folders(cfg) -> tuple[list[tuple[str, str]], list[tuple[str, int]]]:
    """v3.9.33 — pre-flight validasi folder data SEBELUM pipeline jalan.

    Return (missing, counts):
    - missing: [(key, full_path)] utk folder Folders[key] yang path-nya
      TIDAK ada (RootPath sendiri invalid → label 'RootPath').
    - counts: [(key, jumlah file .xls*)] per key Folders (folder hilang → 0).

    Audit 48-f HIGH: folder data hilang / RootPath salah dulu membuat
    pipeline 'SELESAI' hijau + tabel kosong SENYAP. app.py memanggil ini
    sebelum run_pipeline utk menampilkan peringatan eksplisit ke user
    (folder mana yang hilang + berapa file yang akan terbaca).
    """
    root = str(getattr(cfg, "RootPath", "") or "")
    missing: list = []
    counts: list = []
    if not os.path.isdir(root):
        # RootPath sendiri invalid → semua subfolder ikut hilang; catat
        # sebagai entri khusus 'RootPath' di depan.
        missing.append(("RootPath", root))
    folders = getattr(cfg, "Folders", None) or {}
    for key in sorted(folders):
        try:
            full = cfg.folder(key)
        except Exception:  # noqa: BLE001 — cfg duck-typed/mock: fallback join
            full = os.path.join(root, str(folders[key]))
        # hitung file utk SEMUA key (folder hilang → 0, kontrak)
        counts.append((key, len(_list_excel_files(full))))
        if not os.path.isdir(full):
            missing.append((key, full))
    return missing, counts


# =============================================================================
# COLUMN OPERATIONS (mirror M Table functions)
# =============================================================================

def table_select_columns(df: pd.DataFrame, columns: List[str]) -> pd.DataFrame:
    """Mirror M: Table.SelectColumns(table, columns, MissingField.Ignore)."""
    available = [c for c in columns if c in df.columns]
    return df[available].copy()


def table_rename_columns(df: pd.DataFrame, rename_map: dict) -> pd.DataFrame:
    """Mirror M: Table.RenameColumns(table, {{old, new}, ...}, MissingField.Ignore)."""
    valid = {old: new for old, new in rename_map.items() if old in df.columns}
    return df.rename(columns=valid)


def safe_to_numeric(series: pd.Series) -> pd.Series:
    """Convert series ke numeric, null untuk yang gagal. Mirror Table.TransformColumnTypes ke number."""
    return pd.to_numeric(series, errors="coerce")


def safe_to_str(series: pd.Series) -> pd.Series:
    """Convert series ke string, null tetap null. Mirror Table.TransformColumnTypes ke text.
    OPTIMIZED: vectorized, no .apply()."""
    # Use vectorized approach: where notna, convert to str; where na, keep None
    mask = series.notna()
    result = pd.Series(index=series.index, dtype=object)
    result[mask] = series[mask].astype(str)
    result[~mask] = None
    return result


def to_datetime_robust(series: pd.Series) -> pd.Series:
    """
    v3.9.28 — pd.to_datetime yang TAHAN SERIAL EXCEL numerik.

    Sumber masalah: file .xlsb dibaca lewat engine pyxlsb menghasilkan sel
    tanggal sebagai SERIAL MENTAH (mis. 46290 = 2026-09-25) — pyxlsb tidak
    mengonversi sel tanggal ke datetime. pd.to_datetime(46290) membacanya
    sebagai NANODETIK-EPOCH (1970-01-01 + 46µs) → SEMUA nilai jatuh ke
    hari 1 → In_H = 0 untuk H > 1 → 'Stok Seharusnya' salah & flag
    'Hitung Manual' palsu (reproduksi: sheet '4  42 05' workbook user —
    4.591 baris serial → 100% hari 1 pra-fix).

    Aturan: nilai NUMERIC dalam rentang serial Excel wajar (20000..80000
    ≈ th 1954..2119 — tak mungkin tumpang tindih dgn nomor hari 1..31)
    dikonversi via epoch Excel 1899-12-30 (pecahan waktu ikut terbawa);
    nilai lain (datetime/Timestamp/string tanggal) diparse
    pd.to_datetime seperti biasa. Kolom datetime64 apa adanya.
    """
    s = series if isinstance(series, pd.Series) else pd.Series(series)
    if pd.api.types.is_datetime64_any_dtype(s):
        return s
    num = pd.to_numeric(s, errors="coerce")
    serial = num.notna() & num.between(20000.0, 80000.0)
    if not bool(serial.any()):
        return pd.to_datetime(s, errors="coerce")
    # serial Excel (float/int; hasil pyxlsb .xlsb / export numerik) → Timestamp
    ser_ts = pd.Timestamp("1899-12-30") + pd.to_timedelta(
        num.where(serial), unit="D")
    # nilai non-serial → jalur parse biasa (string/datetime objek)
    rest = pd.to_datetime(s.where(~serial), errors="coerce")
    return ser_ts.where(serial, rest)


def resto_from_cabang_po(val) -> str:
    """v3.9.30 — Cabang PO Rincian 32.43 → kode resto: bagian setelah
    titik TERAKHIR ("1370.MLGPAR" → "MLGPAR"; tanpa titik → nilai apa
    adanya). Satu logika bersama pipeline (fn_audit PO Menunggu) + UI
    (rincian_po_view._resto_name delegasi ke sini).

    NA-safety (mirror _doc_code_str — dulu `str(val or "")` mengevaluasi
    bool(pd.NA) → TypeError) + float integral di-stringify TANPA ".0"
    (1370.0 → "1370"; dulu "1370.0" → rsplit titik terakhir → "0")."""
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    s = str(val or "").strip()
    if not s or s.lower() in ("none", "nan"):
        return ""
    # float non-integral (mis. 1370.5): perilaku lama dipertahankan —
    # rsplit titik terakhir ("1370.5" → "5")
    return s.rsplit(".", 1)[-1].strip() if "." in s else s


def date_to_day(series: pd.Series) -> pd.Series:
    """
    Mirror M: Table.TransformColumns(t, {{"Tanggal", Date.Day, Int64.Type}}).
    Ambil day-of-month dari datetime/date column. Return Int64.
    OPTIMIZED: vectorized via .dt accessor.
    v3.9.28: dibangun di atas to_datetime_robust — kolom tanggal numerik
    (serial Excel, mis. dari file .xlsb via pyxlsb) dikonversi benar,
    bukan dibaca sbg nanodetik-epoch → hari 1 (lihat docstring helper).
    """
    dt = to_datetime_robust(series)
    return dt.dt.day.astype("Int64")
